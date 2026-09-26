"""报告产出与健壮性测试。

两层关注点：
1. **正常路径** —— 仓库自带的虚构 demo 必须始终能跑通，且输出结构与文档一致。
2. **异常路径** —— 坏数据、空数据、缺参数都必须"干净失败"：给出可读的
   VERIFY_FAIL / WARN 与退出码，绝不抛 traceback 击穿整单。

第 2 点对应一条真实历史事故：某条线索的 `company` 字段是字符串时，
旧版校验放行，随后在复算阶段抛 AttributeError，整份报告作废。
"""

from __future__ import annotations

import json
import subprocess
import sys

from conftest import (
    ALL_SHEETS,
    DEFAULT_META,
    DEMO_LEADS,
    DEMO_META,
    SHEET_META,
    SHEET_OVERVIEW,
    SHEET_TIER_BC,
    SCRIPT,
    read_demo_leads,
)


# ─────────────────────────────────────────────────────────────
# 正常路径：仓库自带 demo 回归
# ─────────────────────────────────────────────────────────────


def test_demo_regression_is_clean(run_report):
    """demo 数据必须跑出 SUCCESS、A1/B2/C1，且零 WARN。

    demo 是用户 clone 后第一条命令就会跑的东西，也是评分公式的活文档：
    它的分级分布变了，说明公式或数据动了，必须是有意的。
    """
    run = run_report(read_demo_leads(), meta=json.loads(DEMO_META.read_text(encoding="utf-8")))
    assert run.ok, f"demo 未跑通\nstdout={run.stdout}\nstderr={run.stderr}"
    assert "A 1 / B 2 / C 1" in run.stdout
    assert run.warn_lines == [], f"demo 出现意外告警：{run.warn_lines}"


def test_sheet_names_and_order(run_report, lead_factory):
    run = run_report([lead_factory("synthetic-sheets")])
    assert run.sheet_names == ALL_SHEETS


def test_overview_row_count_matches_leads(run_report, lead_factory):
    leads = [lead_factory(f"synthetic-{i}") for i in range(3)]
    run = run_report(leads)
    assert len(run.overview_by_name()) == 3


def test_a_sheet_rows_equal_overview_tier_a_rows(run_report, lead_factory):
    """A 级 Sheet 与总览的 A 计数必须自洽（round-trip 自验的核心断言之一）。"""
    leads = [
        lead_factory("synthetic-a"),
        lead_factory("synthetic-c", cross=1),
    ]
    run = run_report(leads)
    overview_a = {n for n, row in run.overview_by_name().items() if row["tier"] == "A"}
    assert run.tier_a_names() == overview_a


def test_c_tier_rows_carry_next_verification_step(run_report, lead_factory):
    """C 级必须有"下一步验证动作"，不伪装成可触达。"""
    run = run_report([lead_factory("synthetic-needs-step", cross=1)])
    from openpyxl import load_workbook

    ws = load_workbook(run.output)[SHEET_TIER_BC]
    headers = [c.value for c in ws[1]]
    col = headers.index("下一步验证") + 1
    assert ws.max_row >= 2, "该线索应出现在 B-C 名单里"
    assert ws.cell(2, col).value, "C 级线索缺下一步验证动作"


def test_report_contains_disclaimer(run_report, lead_factory):
    """元数据 Sheet 必须含免责声明（诚实边界的硬约束）。"""
    run = run_report([lead_factory("synthetic-disclaimer")])
    from openpyxl import load_workbook

    ws = load_workbook(run.output)[SHEET_META]
    values = [ws.cell(r, 2).value for r in range(1, ws.max_row + 1)]
    assert any(v and "不含海关采购记录" in str(v) for v in values)


# ─────────────────────────────────────────────────────────────
# 异常路径：坏数据必须干净失败
# ─────────────────────────────────────────────────────────────


def test_bad_jsonl_line_is_skipped_with_warning(run_report_raw, lead_factory):
    """一行坏 JSON 只跳过该行，其余线索照常出报告。"""
    good = json.dumps(lead_factory("synthetic-good-line"), ensure_ascii=False)
    run = run_report_raw([good, "{ this is not valid json }"])
    assert run.returncode == 0, run.stderr
    assert "SUCCESS" in run.stdout
    assert any("JSON 解析失败" in line for line in run.warn_lines)
    assert len(run.overview_by_name()) == 1


def test_malformed_lead_type_is_skipped_not_crash(run_report_raw, lead_factory):
    """`company` 为字符串时必须跳过该行 —— 历史事故的回归测试。"""
    broken = dict(lead_factory("synthetic-broken-company"))
    broken["company"] = "Acme Ltd"  # 应为对象
    good = json.dumps(lead_factory("synthetic-survivor"), ensure_ascii=False)

    run = run_report_raw([json.dumps(broken, ensure_ascii=False), good])
    assert run.returncode == 0, f"坏行击穿了整单运行\n{run.stderr}"
    assert "Traceback" not in run.stderr
    assert any("company" in line for line in run.warn_lines)
    assert "Synthetic synthetic-survivor" in run.overview_by_name()


def test_all_leads_malformed_fails_cleanly(run_report_raw):
    """全部线索无效时：VERIFY_FAIL + 退出码 1，而不是 traceback。"""
    run = run_report_raw(['{"company": "not an object"}'])
    assert run.returncode == 1
    assert "VERIFY_FAIL" in run.stderr
    assert "Traceback" not in run.stderr


def test_empty_leads_fails_cleanly(run_report_raw):
    run = run_report_raw([""])
    assert run.returncode == 1
    assert "VERIFY_FAIL" in run.stderr
    assert "Traceback" not in run.stderr


def test_missing_input_argument_fails_cleanly(tmp_path):
    out = tmp_path / "out.xlsx"
    proc = subprocess.run(
        [sys.executable, str(SCRIPT), "--output", str(out)],
        capture_output=True,
        text=True,
    )
    assert proc.returncode == 1
    assert "VERIFY_FAIL" in proc.stderr
    assert "Traceback" not in proc.stderr


# ─────────────────────────────────────────────────────────────
# 备选输入方式：--data / --meta-data
# ─────────────────────────────────────────────────────────────


def test_data_json_string_input_works(tmp_path, lead_factory):
    """`--data` 走 JSON 数组字符串入口，应与 --leads 等价。"""
    payload = json.dumps([lead_factory("synthetic-via-data")], ensure_ascii=False)
    out = tmp_path / "via_data.xlsx"
    proc = subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "--data",
            payload,
            "--meta-data",
            json.dumps(DEFAULT_META, ensure_ascii=False),
            "--output",
            str(out),
        ],
        capture_output=True,
        text=True,
    )
    assert proc.returncode == 0, proc.stderr
    assert "SUCCESS" in proc.stdout

    from openpyxl import load_workbook

    wb = load_workbook(out)
    names = {row[0] for row in wb[SHEET_OVERVIEW].iter_rows(min_row=2, values_only=True) if row and row[0]}
    assert "Synthetic synthetic-via-data" in names


def test_meta_is_optional(tmp_path, lead_factory):
    """不传 meta 时应使用空 meta 正常出报告，而不是报错。"""
    leads_file = tmp_path / "leads.jsonl"
    with leads_file.open("w", encoding="utf-8") as fh:
        fh.write(json.dumps(lead_factory("synthetic-no-meta"), ensure_ascii=False) + "\n")
    out = tmp_path / "no_meta.xlsx"
    proc = subprocess.run(
        [sys.executable, str(SCRIPT), "--leads", str(leads_file), "--output", str(out)],
        capture_output=True,
        text=True,
    )
    assert proc.returncode == 0, proc.stderr
    assert out.exists()


def test_demo_paths_exist():
    """demo 文件必须随仓库分发 —— README 里让用户第一条命令就跑它。"""
    assert DEMO_LEADS.exists()
    assert DEMO_META.exists()
