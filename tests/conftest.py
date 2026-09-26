"""测试共享夹具。

**数据原则**：所有测试数据都在运行期构造，域名只用 RFC 2606 保留域
（`.invalid` 为主，需要测 URL 解析时用 `example.com`），国家/市场标记统一带
`TESTONLY`。不使用任何真实客户、供应商、报价或成交数据 —— 见 `PUBLICATION_POLICY.md`。

**被测对象**：`market-intel/scripts/build_report.py`，通过两种方式验证：

1. **单元级** —— 直接导入脚本模块（该脚本无导入期副作用），测纯函数与
   `_recompute_credibility` 的判定逻辑。
2. **端到端** —— 用 subprocess 像用户那样调用 CLI，再打开生成的 xlsx 核对
   Sheet 内容。这一层才能验证"虚报 A 级到底有没有进 A 级名单"这类结果性断言。
"""

from __future__ import annotations

import importlib.util
import json
import subprocess
import sys
from pathlib import Path

import pytest

REPO_ROOT = Path(__file__).resolve().parent.parent
SKILL_DIR = REPO_ROOT / "market-intel"
SCRIPT = SKILL_DIR / "scripts" / "build_report.py"
DEMO_LEADS = SKILL_DIR / "examples" / "demo_leads.jsonl"
DEMO_META = SKILL_DIR / "examples" / "demo_meta.json"
RUBRIC = SKILL_DIR / "references" / "scoring-rubric.md"
SKILL_MD = SKILL_DIR / "SKILL.md"
SKILL_README = SKILL_DIR / "README.md"

# 仓库交付的三个技能目录（用于跨技能的一致性校验）
SKILLS_DIRS = [
    SKILL_DIR,
    REPO_ROOT / "rfq-to-quotation",
    REPO_ROOT / "verify-docs",
]

# 报告 Sheet 名。test_docs_consistency 会核对 README 里的说法与之一致。
SHEET_OVERVIEW = "线索总览"
SHEET_TIER_A = "A级·可直接发信"
SHEET_TIER_BC = "B-C级·待验证"
SHEET_INTEL = "母语渠道情报"
SHEET_META = "运行元数据"
ALL_SHEETS = [SHEET_OVERVIEW, SHEET_TIER_A, SHEET_TIER_BC, SHEET_INTEL, SHEET_META]

DEFAULT_META = {
    "keywords": ["synthetic"],
    "hs_codes": ["7331.00"],
    "markets": ["TESTONLY"],
    "funnel": {"discovered": 1, "extracted": 1, "verified": 1},
    "cost": {"firecrawl_credits": 0},
    "sources_used": ["synthetic fixture — no real data"],
}


@pytest.fixture(scope="session")
def br():
    """把被测脚本作为模块导入，供单元测试直接调用其内部函数。"""
    spec = importlib.util.spec_from_file_location("build_report_under_test", SCRIPT)
    assert spec is not None and spec.loader is not None, f"无法加载 {SCRIPT}"
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def make_lead(
    lead_id: str,
    *,
    name: str | None = None,
    cross: int = 3,
    decision_makers: list[dict] | None = None,
    email_level: str = "E1",
    homepage_domain: str | None = None,
    email_domain: str | None = None,
    red_flags: list[str] | None = None,
    red_flags_critical: list[str] | None = None,
    claim_tier: str = "A",
    claim_score: float = 0.95,
    subscores: dict | None = None,
) -> dict:
    """构造一条合成线索。

    默认值刻意给一组「看起来完全健康」的参数，测试只在关心的那个维度上做改动，
    这样断言失败时能直接指向被改的维度。
    """
    domain = homepage_domain or f"{lead_id}.example.invalid"
    mail_domain = email_domain or domain
    if decision_makers is None:
        decision_makers = [
            {
                "name": "Synthetic Buyer",
                "role": "Procurement Manager",
                "email": f"purchasing@{mail_domain}",
                "email_level": email_level,
                "source": "synthetic fixture",
                "reachable": True,
            }
        ]

    credibility = {
        "score": claim_score,
        "tier": claim_tier,
        "email_best_level": email_level,
        "red_flags": list(red_flags or []),
        "red_flags_critical": list(red_flags_critical or []),
        "next_verification_step": None,
    }
    if subscores is not None:
        credibility["subscores"] = subscores

    return {
        "lead_id": lead_id,
        "company": {
            "name_local": name or f"Synthetic {lead_id}",
            "name_en": name or f"Synthetic {lead_id}",
            "legal_name_verified": True,
            "country": "TESTONLY",
            "city": "Testville",
            "homepage": f"https://{domain}",
            "product_line_match": ["DIN125 washer"],
            "target_entity": "Synthetic Exporter Co",
        },
        "scale_signals": {
            "employee_hint": "20-50",
            "years_in_business": "10+",
            "branches": 2,
            "is_distributor": True,
            "size_confidence": "high",
        },
        "decision_makers": decision_makers,
        "lead_source": {
            "discovery_channel": "synthetic fixture",
            "discovery_query": "synthetic",
            "cross_source_count": cross,
        },
        "credibility": credibility,
        "bd_strategy": {
            "entry_angle": "synthetic",
            "recommended_channel": "email",
            "template_hint": "synthetic",
            "priority": "high",
        },
    }


@pytest.fixture
def lead_factory():
    """合成线索工厂，返回 `make_lead`。所有测试数据均在运行期生成。"""
    return make_lead


class ReportRun:
    """一次端到端运行的结果，带读取生成 Excel 的便捷方法。"""

    def __init__(self, proc: subprocess.CompletedProcess, output: Path, leads: list[dict]):
        self.proc = proc
        self.returncode = proc.returncode
        self.stdout = proc.stdout
        self.stderr = proc.stderr
        self.output = output
        self.leads = leads

    @property
    def ok(self) -> bool:
        return self.returncode == 0 and "SUCCESS" in self.stdout

    @property
    def warn_lines(self) -> list[str]:
        return [ln for ln in self.stderr.splitlines() if ln.strip()]

    def _workbook(self):
        from openpyxl import load_workbook

        return load_workbook(self.output)

    @property
    def sheet_names(self) -> list[str]:
        return self._workbook().sheetnames

    def tier_a_names(self) -> set[str]:
        """「A级·可直接发信」名单里的公司名集合。"""
        wb = self._workbook()
        ws = wb[SHEET_TIER_A]
        return {str(row[0]) for row in ws.iter_rows(min_row=2, values_only=True) if row and row[0]}

    def tier_bc_names(self) -> set[str]:
        wb = self._workbook()
        ws = wb[SHEET_TIER_BC]
        return {str(row[0]) for row in ws.iter_rows(min_row=2, values_only=True) if row and row[0]}

    def overview_by_name(self) -> dict[str, dict]:
        """「线索总览」按公司名索引：{'tier': ..., 'score': ..., 'email': ...}。"""
        wb = self._workbook()
        ws = wb[SHEET_OVERVIEW]
        headers = [c.value for c in ws[1]]
        out: dict[str, dict] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            out[str(row[0])] = dict(zip(headers, row))
        return out


@pytest.fixture
def run_report(tmp_path):
    """端到端跑一次 `build_report.py`（subprocess，与用户调用方式一致）。"""
    counter = {"n": 0}

    def _run(
        leads: list[dict],
        *,
        meta: dict | None = None,
        expect_success: bool = True,
    ) -> ReportRun:
        counter["n"] += 1
        base = tmp_path / f"run{counter['n']}"
        base.mkdir()

        leads_file = base / "leads.jsonl"
        with leads_file.open("w", encoding="utf-8") as fh:
            for lead in leads:
                fh.write(json.dumps(lead, ensure_ascii=False) + "\n")

        meta_file = base / "meta.json"
        meta_file.write_text(
            json.dumps(meta if meta is not None else DEFAULT_META, ensure_ascii=False),
            encoding="utf-8",
        )

        output = base / "report.xlsx"
        proc = subprocess.run(
            [
                sys.executable,
                str(SCRIPT),
                "--leads",
                str(leads_file),
                "--meta",
                str(meta_file),
                "--output",
                str(output),
            ],
            capture_output=True,
            text=True,
        )
        if expect_success and proc.returncode != 0:
            pytest.fail(
                f"报告生成失败（exit {proc.returncode}）\n"
                f"stdout:\n{proc.stdout}\nstderr:\n{proc.stderr}"
            )
        return ReportRun(proc, output, leads)

    return _run


@pytest.fixture
def run_report_raw(tmp_path):
    """同 run_report，但不会因失败而 fail —— 用于测错误路径与坏数据。"""
    counter = {"n": 0}

    def _run(lines: list[str], *, expect_file: bool = True) -> ReportRun:
        counter["n"] += 1
        base = tmp_path / f"raw{counter['n']}"
        base.mkdir()
        leads_file = base / "leads.jsonl"
        leads_file.write_text("\n".join(lines) + "\n", encoding="utf-8")
        meta_file = base / "meta.json"
        meta_file.write_text(json.dumps(DEFAULT_META, ensure_ascii=False), encoding="utf-8")
        output = base / "report.xlsx"
        proc = subprocess.run(
            [
                sys.executable,
                str(SCRIPT),
                "--leads",
                str(leads_file),
                "--meta",
                str(meta_file),
                "--output",
                str(output),
            ],
            capture_output=True,
            text=True,
        )
        return ReportRun(proc, output, [])

    return _run


def read_demo_leads() -> list[dict]:
    return [json.loads(ln) for ln in DEMO_LEADS.read_text(encoding="utf-8").splitlines() if ln.strip()]


def strip_code_fences(text: str) -> str:
    """去掉围栏代码块，避免把代码示例里的转义序列当成正文排版错误。"""
    out, inside = [], False
    for line in text.splitlines():
        if line.lstrip().startswith("```"):
            inside = not inside
            continue
        if not inside:
            out.append(line)
    return "\n".join(out)
